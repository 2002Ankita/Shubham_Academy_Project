from fastapi import APIRouter, Depends, HTTPException, Query, Response
from typing import List, Optional
from datetime import datetime
import io
from bson import ObjectId

from app.models.leave_request import LeaveRequest
from app.schemas.leave_request import LeaveRequestCreate, LeaveRequestResponse
from app.services.auth_service import get_current_user
from app.services.teacher_service import resolve_teacher
from app.routers.auth import oauth2_scheme

router = APIRouter(prefix="/leave-requests", tags=["Leave Requests"])


def _extract_teacher_id(req: LeaveRequest) -> str:
    """Safely extract teacher ObjectId as string from LeaveRequest document."""
    if hasattr(req, "teacher") and req.teacher:
        if hasattr(req.teacher, "id"):
            return str(req.teacher.id)
        if hasattr(req.teacher, "ref") and hasattr(req.teacher.ref, "id"):
            return str(req.teacher.ref.id)
    return ""


def _generate_leave_excel(requests: list, teacher_name: str = None, balance_summary: dict = None) -> bytes:
    """Generate an executive-styled .xlsx workbook for faculty leave requests."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="Excel generation requires openpyxl. Please run: pip install openpyxl"
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leave Requests"

    # Styling definitions
    title_font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="8B1216", end_color="8B1216", fill_type="solid")

    meta_label_font = Font(name="Segoe UI", size=10, bold=True, color="334155")
    meta_val_font = Font(name="Segoe UI", size=10, color="0F172A")

    header_font = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    regular_font = Font(name="Segoe UI", size=10, color="0F172A")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # 1. Title Banner
    ws.merge_cells("A1:K1")
    t_cell = ws["A1"]
    t_cell.value = "SHUBHAM ACADEMY — FACULTY LEAVE REQUESTS REPORT"
    t_cell.font = title_font
    t_cell.fill = title_fill
    t_cell.alignment = center_align
    ws.row_dimensions[1].height = 36

    # 2. Metadata Block
    current_time_str = datetime.now().strftime("%d %b %Y, %I:%M %p")
    t_name = teacher_name or "All Faculty Members"

    ws["A2"] = "Faculty Member:"
    ws["B2"] = t_name
    ws["D2"] = "Report Generated:"
    ws["E2"] = current_time_str

    ws["A2"].font = meta_label_font
    ws["B2"].font = meta_val_font
    ws["D2"].font = meta_label_font
    ws["E2"].font = meta_val_font

    total_entitled = balance_summary.get("total", 24) if balance_summary else 24
    approved_taken = balance_summary.get("taken", sum(r.number_of_days for r in requests if r.status == "Approved")) if balance_summary else sum(r.number_of_days for r in requests if r.status == "Approved")
    pending_count = balance_summary.get("pending", sum(1 for r in requests if r.status == "Pending")) if balance_summary else sum(1 for r in requests if r.status == "Pending")
    available_leaves = balance_summary.get("available", total_entitled - approved_taken) if balance_summary else total_entitled - approved_taken

    ws["A3"] = "Total Entitled:"
    ws["B3"] = f"{total_entitled} Days"
    ws["D3"] = "Leaves Taken:"
    ws["E3"] = f"{approved_taken} Days"
    ws["G3"] = "Pending Approval:"
    ws["H3"] = f"{pending_count} Request(s)"
    ws["J3"] = "Available Balance:"
    ws["K3"] = f"{available_leaves} Days"

    for c in ["A3", "D3", "G3", "J3"]:
        ws[c].font = meta_label_font
    for c in ["B3", "E3", "H3", "K3"]:
        ws[c].font = meta_val_font

    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 10

    # 3. Table Column Headers
    headers = [
        "Sr No",
        "Request ID",
        "Faculty Name",
        "Leave Type",
        "From Date",
        "To Date",
        "Days",
        "Reason / Purpose",
        "Approval Status",
        "Applied Date",
        "Admin Remarks"
    ]

    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[5].height = 28

    # 4. Table Data Rows
    row_num = 6
    if not requests:
        ws.merge_cells("A6:K6")
        empty_cell = ws["A6"]
        empty_cell.value = "No leave request records found."
        empty_cell.font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        empty_cell.alignment = center_align
        ws.row_dimensions[6].height = 26
    else:
        for idx, r in enumerate(requests, start=1):
            is_even = (idx % 2 == 0)
            row_fill = PatternFill(start_color="F8FAFC" if is_even else "FFFFFF", fill_type="solid")

            req_id = str(r.id)
            fac_name = r.teacher_name or t_name
            l_type = r.leave_type or "Leave"
            from_str = r.from_date.strftime("%d-%b-%Y") if r.from_date else "-"
            to_str = r.to_date.strftime("%d-%b-%Y") if r.to_date else "-"
            days_val = r.number_of_days or 0
            reason_str = r.reason or "-"
            status_val = r.status or "Pending"
            created_str = r.created_at.strftime("%d-%b-%Y") if r.created_at else "-"
            comment_str = r.admin_comment or "-"

            row_values = [
                idx,
                req_id,
                fac_name,
                l_type,
                from_str,
                to_str,
                days_val,
                reason_str,
                status_val,
                created_str,
                comment_str
            ]

            for col_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=row_num, column=col_idx, value=val)
                cell.font = regular_font
                cell.fill = row_fill
                cell.border = thin_border

                if col_idx in [1, 7]:
                    cell.alignment = center_align
                elif col_idx in [5, 6, 10]:
                    cell.alignment = center_align
                elif col_idx == 9:
                    cell.alignment = center_align
                    if status_val == "Approved":
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="166534")
                        cell.fill = PatternFill(start_color="DCFCE7", fill_type="solid")
                    elif status_val == "Rejected":
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")
                        cell.fill = PatternFill(start_color="FEE2E2", fill_type="solid")
                    elif status_val == "Cancelled":
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="475569")
                        cell.fill = PatternFill(start_color="F1F5F9", fill_type="solid")
                    else:
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="92400E")
                        cell.fill = PatternFill(start_color="FEF3C7", fill_type="solid")
                else:
                    cell.alignment = left_align

            ws.row_dimensions[row_num].height = 24
            row_num += 1

    # Auto column width calculation
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2, 3, 4]:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.post("", response_model=LeaveRequestResponse)
async def create_leave_request(
    request_in: LeaveRequestCreate,
    token: str = Depends(oauth2_scheme)
):
    user = await get_current_user(token)
    
    # Resolve teacher by request_in.teacher_id or user.id
    target_id = request_in.teacher_id or str(user.id)
    if target_id.lower() == "me":
        target_id = str(user.id)
        
    teacher = await resolve_teacher(target_id)
    if not teacher and str(user.id) != target_id:
        teacher = await resolve_teacher(str(user.id))
    if not teacher:
        from app.models.teacher import Teacher
        teacher = await Teacher.find_one()
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher profile not found")
    
    delta = request_in.to_date - request_in.from_date
    days = delta.days + 1
    if days < 1:
        raise HTTPException(status_code=400, detail="Invalid date range: To Date cannot be earlier than From Date")

    # Name resolution
    teacher_name = (
        getattr(request_in, "teacher_name", None)
        or ((teacher.user.full_name if hasattr(teacher.user, "full_name") else None) if teacher else None)
        or user.full_name
        or "Faculty Member"
    )

    # Branch resolution
    from app.core.branch_config import normalize_branch, BRANCH_MANGALWAR
    target_branch = (
        getattr(request_in, "branch", None)
        or getattr(teacher, "branch", None)
        or getattr(user, "campus", None)
        or getattr(user, "branch", None)
        or BRANCH_MANGALWAR
    )
    target_branch = normalize_branch(target_branch) or target_branch or BRANCH_MANGALWAR

    leave_req = LeaveRequest(
        teacher=teacher,
        teacher_name=teacher_name,
        leave_type=request_in.leave_type,
        from_date=request_in.from_date,
        to_date=request_in.to_date,
        number_of_days=days,
        branch=target_branch,
        reason=request_in.reason,
        status="Pending",
        created_at=datetime.utcnow()
    )
    await leave_req.insert()
    
    return LeaveRequestResponse(
        id=str(leave_req.id),
        teacher_id=str(teacher.id),
        teacher_name=leave_req.teacher_name,
        leave_type=leave_req.leave_type,
        from_date=leave_req.from_date,
        to_date=leave_req.to_date,
        number_of_days=leave_req.number_of_days,
        branch=leave_req.branch,
        reason=leave_req.reason,
        status=leave_req.status,
        admin_comment=leave_req.admin_comment,
        created_at=leave_req.created_at,
        branch=leave_req.branch
    )


@router.get("/export")
async def export_leave_requests_excel(
    teacher_id: Optional[str] = None,
    branch: Optional[str] = None,
    token: str = Depends(oauth2_scheme)
):
    """Export leave requests to an executive Excel spreadsheet."""
    user = await get_current_user(token)
    await ensure_default_leave_requests()
    
    teacher_name = None
    balance_summary = None

    from app.core.branch_config import normalize_branch

    if teacher_id and teacher_id.lower() != "all":
        t_id = str(user.id) if teacher_id.lower() == "me" else teacher_id
        teacher = await resolve_teacher(t_id)
        if not teacher:
            raise HTTPException(status_code=404, detail="Teacher profile not found")
        requests = await LeaveRequest.find(LeaveRequest.teacher.id == teacher.id).sort("-created_at").to_list()
        
        # User name resolution
        if hasattr(teacher, "user") and teacher.user and hasattr(teacher.user, "full_name"):
            teacher_name = teacher.user.full_name
        else:
            teacher_name = user.full_name or "Faculty Member"

        taken = sum(r.number_of_days for r in requests if r.status == "Approved")
        pending = sum(1 for r in requests if r.status == "Pending")
        total = 24
        balance_summary = {
            "total": total,
            "taken": taken,
            "pending": pending,
            "available": total - taken
        }
    else:
        all_reqs = await LeaveRequest.find_all().sort("-created_at").to_list()
        if branch and branch.lower() != "all":
            norm_b = normalize_branch(branch)
            requests = [r for r in all_reqs if normalize_branch(getattr(r, "branch", None)) == norm_b]
            teacher_name = f"All Faculty ({norm_b})"
        else:
            requests = all_reqs
            teacher_name = "All Faculty"

    excel_bytes = _generate_leave_excel(requests, teacher_name=teacher_name, balance_summary=balance_summary)
    
    clean_name = (teacher_name or "Faculty").replace(" ", "_").replace(".", "").replace("(", "").replace(")", "")
    filename = f"Leave_Requests_{clean_name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    
    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )


@router.get("/teacher/{teacher_id}/export")
async def export_teacher_leave_requests_excel(
    teacher_id: str,
    token: str = Depends(oauth2_scheme)
):
    """Convenience endpoint to export a specific teacher's leave history to Excel."""
    return await export_leave_requests_excel(teacher_id=teacher_id, token=token)


@router.get("/teacher/{teacher_id}/summary")
async def get_teacher_leave_summary(teacher_id: str, token: str = Depends(oauth2_scheme)):
    """Fetch aggregated leave balance stats for a specific faculty member."""
    user = await get_current_user(token)
    t_id = str(user.id) if teacher_id.lower() == "me" else teacher_id
    teacher = await resolve_teacher(t_id)
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher profile not found")

    requests = await LeaveRequest.find(LeaveRequest.teacher.id == teacher.id).to_list()
    total = 24
    taken = sum(r.number_of_days for r in requests if r.status == "Approved")
    pending = sum(1 for r in requests if r.status == "Pending")
    available = max(0, total - taken)

    return {
        "teacher_id": str(teacher.id),
        "total": total,
        "taken": taken,
        "pending": pending,
        "available": available
    }


@router.get("/teacher/{teacher_id}", response_model=List[LeaveRequestResponse])
async def get_teacher_leave_requests(teacher_id: str, token: str = Depends(oauth2_scheme)):
    user = await get_current_user(token)
    t_id = str(user.id) if teacher_id.lower() == "me" else teacher_id
    teacher = await resolve_teacher(t_id)
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher profile not found")
        
    requests = await LeaveRequest.find(LeaveRequest.teacher.id == teacher.id).sort("-created_at").to_list()
    return [
        LeaveRequestResponse(
            id=str(r.id),
            teacher_id=str(teacher.id),
            teacher_name=r.teacher_name,
            leave_type=r.leave_type,
            from_date=r.from_date,
            to_date=r.to_date,
            number_of_days=r.number_of_days,
            branch=getattr(r, "branch", None) or getattr(teacher, "branch", None) or "Mangalvar Peth",
            reason=r.reason,
            status=r.status,
            admin_comment=r.admin_comment,
            created_at=r.created_at,
            branch=getattr(r, "branch", getattr(teacher, "branch", "Tarabai Park"))
        ) for r in requests
    ]


@router.get("", response_model=List[LeaveRequestResponse])
async def get_all_leave_requests(
    branch: Optional[str] = Query(None),
    token: str = Depends(oauth2_scheme)
):
    await get_current_user(token)
    await ensure_default_leave_requests()
    
    requests = await LeaveRequest.find_all().sort("-created_at").to_list()
    from app.core.branch_config import normalize_branch, BRANCH_TARABAI, BRANCH_MANGALWAR
    
    norm_filter = normalize_branch(branch) if branch and branch.lower() != "all" else None
    
    result = []
    for r in requests:
        r_branch = getattr(r, "branch", None)
        if not r_branch and r.teacher:
            try:
                t_obj = r.teacher if hasattr(r.teacher, "branch") else await resolve_teacher(str(r.teacher.ref.id))
                if t_obj and getattr(t_obj, "branch", None):
                    r_branch = t_obj.branch
            except Exception:
                pass
        if not r_branch:
            r_branch = BRANCH_MANGALWAR if "mangal" in (r.teacher_name or "").lower() else BRANCH_TARABAI

        norm_r_branch = normalize_branch(r_branch) or r_branch

        if norm_filter and norm_r_branch != norm_filter:
            continue

        result.append(
            LeaveRequestResponse(
                id=str(r.id),
                teacher_id=_extract_teacher_id(r),
                teacher_name=r.teacher_name,
                leave_type=r.leave_type,
                from_date=r.from_date,
                to_date=r.to_date,
                number_of_days=r.number_of_days,
                reason=r.reason,
                status=r.status,
                admin_comment=r.admin_comment,
                created_at=r.created_at,
                branch=norm_r_branch or "Mangalvar Peth"
            )
        )
    return result


@router.put("/{request_id}/status")
async def update_leave_status(
    request_id: str,
    status: Optional[str] = Query(None),
    admin_comment: Optional[str] = Query(None),
    token: str = Depends(oauth2_scheme)
):
    await get_current_user(token)
    try:
        req = await LeaveRequest.get(ObjectId(request_id))
    except Exception:
        req = None
    if not req:
        raise HTTPException(status_code=404, detail="Leave request not found")
    if status:
        req.status = status
    if admin_comment is not None:
        req.admin_comment = admin_comment
    await req.save()
    return {
        "message": "Leave request updated successfully",
        "id": request_id,
        "status": req.status,
        "admin_comment": req.admin_comment,
        "branch": getattr(req, "branch", None)
    }


async def ensure_default_leave_requests():
    """Ensure baseline leave requests exist for both Mangalwar Peth (MP) and Tarabai Park (TP)."""
    count = await LeaveRequest.count()
    if count >= 6:
        return

    from app.models.teacher import Teacher
    from app.core.branch_config import BRANCH_TARABAI, BRANCH_MANGALWAR
    from datetime import timedelta

    teachers = await Teacher.find_all().to_list()
    teacher_mp = None
    teacher_tp = None

    for t in teachers:
        b = getattr(t, "branch", "") or ""
        if "mangal" in b.lower() and not teacher_mp:
            teacher_mp = t
        if "tarabai" in b.lower() and not teacher_tp:
            teacher_tp = t

    if not teacher_tp and teachers:
        teacher_tp = teachers[0]
    if not teacher_mp and len(teachers) > 1:
        teacher_mp = teachers[1]
    elif not teacher_mp and teachers:
        teacher_mp = teachers[0]

    if not teacher_tp:
        return

    today = datetime.utcnow()

    default_requests = [
        # Mangalwar Peth (MP) Faculty Leaves
        {
            "teacher": teacher_mp or teacher_tp,
            "teacher_name": "Prof. Rajesh Sharma",
            "leave_type": "Casual Leave",
            "from_date": today + timedelta(days=2),
            "to_date": today + timedelta(days=3),
            "number_of_days": 2,
            "reason": "Attending National Mathematics Symposium at Pune University",
            "status": "Pending",
            "branch": BRANCH_MANGALWAR,
            "created_at": today - timedelta(days=1),
            "admin_comment": None
        },
        {
            "teacher": teacher_mp or teacher_tp,
            "teacher_name": "Dr. Priya Kulkarni",
            "leave_type": "Medical Leave",
            "from_date": today + timedelta(days=4),
            "to_date": today + timedelta(days=5),
            "number_of_days": 2,
            "reason": "Dental surgery and doctor advised bed rest",
            "status": "Pending",
            "branch": BRANCH_MANGALWAR,
            "created_at": today - timedelta(days=2),
            "admin_comment": None
        },
        {
            "teacher": teacher_mp or teacher_tp,
            "teacher_name": "Prof. Amit Sawant",
            "leave_type": "Earned Leave",
            "from_date": today - timedelta(days=5),
            "to_date": today - timedelta(days=4),
            "number_of_days": 2,
            "reason": "Personal family engagement in Sangli",
            "status": "Approved",
            "branch": BRANCH_MANGALWAR,
            "created_at": today - timedelta(days=7),
            "admin_comment": "Approved by Mangalwar Peth Admin Desk"
        },
        # Tarabai Park (TP) Faculty Leaves
        {
            "teacher": teacher_tp,
            "teacher_name": "Dr. Priya Kulkarni",
            "leave_type": "Casual Leave",
            "from_date": today + timedelta(days=1),
            "to_date": today + timedelta(days=1),
            "number_of_days": 1,
            "reason": "Academic Board Meeting at Shivaji University",
            "status": "Pending",
            "branch": BRANCH_TARABAI,
            "created_at": today - timedelta(days=1),
            "admin_comment": None
        },
        {
            "teacher": teacher_tp,
            "teacher_name": "Prof. Sunita Patil",
            "leave_type": "Sick Leave",
            "from_date": today + timedelta(days=3),
            "to_date": today + timedelta(days=4),
            "number_of_days": 2,
            "reason": "Severe throat infection and vocal rest recommended by ENT specialist",
            "status": "Pending",
            "branch": BRANCH_TARABAI,
            "created_at": today - timedelta(days=2),
            "admin_comment": None
        },
        {
            "teacher": teacher_tp,
            "teacher_name": "Prof. Amit Sawant",
            "leave_type": "Academic Leave",
            "from_date": today - timedelta(days=3),
            "to_date": today - timedelta(days=3),
            "number_of_days": 1,
            "reason": "Delivering guest lecture on Organic Chemistry at Kolhapur College",
            "status": "Approved",
            "branch": BRANCH_TARABAI,
            "created_at": today - timedelta(days=6),
            "admin_comment": "Approved by Tarabai Park Admin Desk"
        }
    ]

    for req_data in default_requests:
        existing = await LeaveRequest.find_one({
            "teacher_name": req_data["teacher_name"],
            "branch": req_data["branch"],
            "leave_type": req_data["leave_type"],
            "reason": req_data["reason"]
        })
        if not existing:
            new_req = LeaveRequest(**req_data)
            await new_req.insert()
